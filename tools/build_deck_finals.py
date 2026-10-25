# Finals deck, live and printed, from one set via run.py --finals-pack.

import re
import shutil
import subprocess
import sys
from contextlib import contextmanager
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
# Repo root on sys.path, so fiam and deck 1's helpers always import.
sys.path.insert(0, str(ROOT))

from fiam import paths  # noqa: E402
from fiam.paths import read_json as load  # noqa: E402
from fiam.report import animation as A  # noqa: E402
from fiam.report import docs as R  # noqa: E402
from fiam.report import qa  # noqa: E402
from tools import build_deck as D  # noqa: E402
from tools.motion import added, finish, stagger  # noqa: E402

OUT, PRINTED = ROOT / "deck" / "FIAM_finals.pptx", ROOT / "deck" / "FIAM_finals_print.pptx"
FIGURES, V2 = ROOT / "out" / "figures", ROOT / "out" / "v2"
CLIPS = FIGURES / "clips"
# The film's palette, so the slides' dots match the film's investors.
LIGHT, NAVY = RGBColor.from_string(A.LIGHT[1:]), RGBColor.from_string(A.NAVY[1:])
PERSONAS = [c[1:] for c in A.PERSONAS]
RECIPES = [A.RECIPES[k][1:] for k in ("v1", "brief_2026", "case_2025")]
CLEAN = "Decisions made January 2025 to July 2026, after the model's training ended"
TEAM = "QUARCC Purple"
PEOPLE = (
    "Nicholas Xydis",
    "Yohannah Ilieva",
    "Andrew Koumoudjian",
    "Nicolas Romero",
    "Nokto Lapointe",
)
MEMBERS = "  ·  ".join(PEOPLE)
SCHOOLS = "Concordia University  ·  McGill University"
# Deck 1's footer, and ours with the team's name in it.
FOOTERS = (
    "FIAM  ·  McGill-FIAM Asset Management Hackathon 2026",
    f"{TEAM}  ·  McGill-FIAM Asset Management Hackathon 2026 Finals",
)
# "live" (the talk) or "print" (the PDF).
MODE = "live"

# The talk slide by slide: what to say, numbers filled from qa.facts().
NOTES = dict(
    title="Open on the idea, not the numbers. In ten minutes, a book that holds up, what it can "
    "carry, and the AI crowd we built to see who we trade against.",
    hook="Your own brief warns that AI agents will crowd into the same names. Rather than guess "
    "what that crowd does, we built it. Sixty AI investors on our own machine, every decision "
    "kept.",
    ideas="Three ideas, one book. Theory picks the stocks. Machine learning keeps the book "
    "neutral. And an AI crowd we built shows us who is on the other side of our trades.",
    journey="One real month, {month}. Of {stocks} US stocks, {tradeable} are big and liquid "
    "enough to trade. Twelve published signals give each one a score, we rank them, and our "
    "optimizer holds the best long and the worst short, {longs} and {shorts} that month, with "
    "zero net bet on the market and on every sector.",
    ml="We use machine learning for risk, not for picking stocks, and only where it beats a "
    "simple rival. It earned its place forecasting beta. It broke on volatility in 2020, so the "
    "simple model runs there.",
    desk="Around the portfolio manager sit six AI colleagues. A rules reader, triage, an analyst, "
    "a devil's advocate, a memo writer and an explainer. They read, check and argue, and none of "
    "them sets a position. On the final book the devil's advocate raised {fv3_asked} critiques "
    "and code confirmed {fv3_confirmed}.",
    results="So, did it work?",
    growth="A dollar invested in January 2021, after every cost, against the hurdle the brief sets "
    "and against the market. It cleared the hurdle with an IR of {ir}.",
    neutral="The brief checks neutrality first. Our beta to the market is {signed_beta}, "
    "statistically zero, and it stays inside the noise band window after window. Its "
    "correlation with the S&P 500 is just {rho}.",
    edge="Every backtest looks good. So is ours real?",
    monkeys="We drew a thousand random long-short books from the five hundred largest stocks, each "
    "with our number of positions, our gross, our holding time and our costs. The final book "
    "beats {null_beaten} of them. The best of them reaches an IR of {null_best}. We make "
    "{fv3_ir_1e7}.",
    ladder="Every benchmark we could think of, on the same money. Equal weights on our score make "
    "{naive_model_10m}. Our favourite names held still, picked with hindsight, make {fav_10m}. "
    "Our rule on the five hundred largest stocks makes {large_10m}. After the haircuts published "
    "research applies to every backtest, we expect an IR of {live_low} to {live_high} live, still "
    "above the hurdle.",
    capacity="The question every PM asks. How much money can it run? Our first book earns the "
    "most near {v1_peak} and stops paying around {v1_breakeven}. The final book still pays at "
    "five billion, with an IR of {fv3_ir_5e9}.",
    final="So this is the book we present. Our first book's signals, traded the way a big fund "
    "must. It slows its trading as the fund grows, holds beta near zero, aims its risk at ten "
    "percent a year and brakes after a fall. Every rule was written down before it ran, and the "
    "ones that failed their own tests were dropped. Against our first book and the two "
    "alternatives it has the best Sharpe at {fd_best_sharpe} of {fd_sizes} fund sizes, and it "
    "still pays at five billion.",
    risk="Where does its risk sit? Each month our risk model splits it name by name. The forecast "
    "is about {risk_vol} a year, {risk_factor_share} of it from styles and the rest stock by "
    "stock, and the ten largest contributors carry {risk_top10}. No single name runs the book.",
    exposures="What is it betting on? Month by month, its tilt to five styles, told by our "
    "explainer agent from that table with every number checked. {factor_story}",
    losses="Weakest first. Before 2021 a small short tilt in a rising market cost us. Since 2023 "
    "the edge has been flat, because a junk rally, when unprofitable stocks fly, is the one thing "
    "this book fears. Replayed through {crises} past crises, today's book made money in "
    "{crises_up} and never lost more than {crisis_worst}.",
    turn="Who is on the other side of our trades? More and more, it is AI.",
    meet="Meet the crowd. Thirty retail investors, six personas each fed different data, and "
    "thirty copycat funds, each handed one paragraph of a strategy, ours among them.",
    swarm="Watch them herd. Thirty AI investors with six different personas, and {biggest} of "
    "them end up voting together. As a group they are worth about {voters} independent "
    "opinions. What they see drives them more than the persona we gave them.",
    film="The film plays by itself. One real month of their decisions, every vote landing, the "
    "herd forming, and our own book among their favourites. Talk over it.",
    copycats="We expected copycats to find our stocks when handed our recipe, and they do. The "
    "surprise is that, handed only the organizers' brief, they land on our stocks even more "
    "often. That is the crowding risk the brief warns about, measured.",
    trades="The point for a PM. Chance would make ten percent of our longs AI crowd favourites, "
    "and {fv3_overlap_retail} are. Our returns load {fv3_crowd_beta} on what the "
    "crowd likes, and we can put a number on what its sell-off would cost us. That is a risk you "
    "can now measure.",
    receipt="Every one of the {decisions_total} decisions is on file. Here is one. What the "
    "copycat was told, the anonymous card it saw, what it answered, and who the codes really "
    "were. Ask it the same card again and it answers the same. Nudge the numbers and it changes "
    "its mind, so it reads the card rather than remembering the stock.",
    today="Today's rules, run live. The final book under the finals mandate, with our first "
    "book's method and the size build beside it, after every cost.",
    launch="Three things we would take to a client. A book that grows with the money, a monitor "
    "for the AI crowd, and a desk that turns new rules into a compliant portfolio in minutes.",
    closing="To close. The right book for the money, the AI crowd watched like any other risk, "
    "and every test written down first. Thank you.",
)


# A slide's notes with its numbers filled in.
def notes(name: str, f: dict, **extra) -> str:
    return NOTES[name].format_map(f | extra)


# A slide's section, or "Appendix" when the printed deck moves it there.
def section(name: str, appendix: bool = False) -> str:
    return "Appendix" if MODE == "print" and appendix else name


# Method pages describe the models, so they read the submitted run.
@contextmanager
def submitted():
    saved = R.SUBMISSION
    R.SUBMISSION = ROOT / "submission"
    try:
        yield
    finally:
        R.SUBMISSION = saved


# Text in a box at (left, top, width, height) in inches, one style.
def say(page, words, at, size=18, color=D.INK, bold=False, center=False):
    box = page.shapes.add_textbox(*(Inches(x) for x in at))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate([words] if isinstance(words, str) else words):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        D.run(para, line, size, color, bold)
    return box


# A flat shape at a spot in inches, named so Morph can match it.
def shape(page, kind, at, fill, name: str = ""):
    s = D.flat(page.shapes.add_shape(kind, *(Inches(x) for x in at)), fill)
    s.name = name or s.name
    return s


# A big number on a light card, label under it, grouped to move as one.
def card(page, at, value: str, label: str, size: float = 32):
    group = page.shapes.add_group_shape()
    shape(group, MSO_SHAPE.RECTANGLE, at, D.PANEL)
    box = say(group, value, (at[0] + 0.25, at[1] + 0.18, at[2] - 0.5, at[3] - 0.3), size, D.ACCENT)
    box.text_frame.paragraphs[0].runs[0].font.bold = True
    D.run(box.text_frame.add_paragraph(), label, 12, D.MUTED)
    return group


# Cards side by side across a width.
def cards(page, items, top=5.3, left=0.6, width=12.13, height=1.25, size=32) -> list:
    each = (width - 0.3 * (len(items) - 1)) / len(items)
    return [
        card(page, (left + (each + 0.3) * i, top, each, height), value, label, size)
        for i, (value, label) in enumerate(items)
    ]


# Cards down a column: left, top, width, height, step down, number size.
def column(page, items, left, top, width, height, step, size) -> list:
    return [
        card(page, (left, top + step * i, width, height), value, label, size)
        for i, (value, label) in enumerate(items)
    ]


# Heading and line on a panel, grouped, to build every visual slide.
def tile(page, at, head: str, line: str = "", fill=D.PANEL, ink=D.INK, size=20):
    group = page.shapes.add_group_shape()
    shape(group, MSO_SHAPE.ROUNDED_RECTANGLE, at, fill)
    say(group, head, (at[0] + 0.25, at[1] + 0.2, at[2] - 0.5, 0.6), size, ink, True)
    if line:
        below = at[1] + 0.3 + size / 40
        say(group, line, (at[0] + 0.25, below, at[2] - 0.5, 1.2), 13 + size // 10, D.MUTED)
    return group


# Numbered panels across the slide, rising one after another.
def pillars(page, ideas) -> list:
    row = []
    for i, (head, line) in enumerate(ideas):
        x = 0.6 + 4.1 * i
        group = page.shapes.add_group_shape()
        shape(group, MSO_SHAPE.ROUNDED_RECTANGLE, (x, 1.85, 3.8, 4.6), D.PANEL)
        say(group, str(i + 1), (x + 0.35, 2.2, 1, 0.9), 54, D.ACCENT, True)
        say(group, head, (x + 0.35, 3.35, 3.2, 1.2), 24, D.INK, True)
        say(group, line, (x + 0.35, 4.75, 3.2, 1.6), 15, D.MUTED)
        row.append(group)
    return stagger(row, "rise", 0, 450)


# A chart fit to its box, from out/figures or, for the book, deck 1's.
def chart(page, name: str, at, book: bool = False):
    path = (D.FIGURES if book else FIGURES) / f"{name}.png"
    if not path.exists():
        return say(page, f"{name} not drawn yet", (at[0], at[1], at[2], 0.4), 11, D.MUTED)
    pic = page.shapes.add_picture(str(path), Inches(at[0]), Inches(at[1]), width=Inches(at[2]))
    if pic.height > Inches(at[3]):
        ratio = Inches(at[3]) / pic.height
        pic.height, pic.width = int(Inches(at[3])), int(pic.width * ratio)
    return pic


# The chart's film (fiam.report.clips) for the timeline, else its still.
def moving(page, name: str, at, still: str, book: bool = False):
    movie, poster = CLIPS / f"{name}.mp4", CLIPS / f"{name}.png"
    if not (movie.exists() and poster.exists()):
        return chart(page, still, at, book), None
    box = (Inches(x) for x in at)
    clip = page.shapes.add_movie(
        str(movie), *box, poster_frame_image=str(poster), mime_type="video/mp4"
    )
    return clip, clip


# Navy slide: small line, headline, its second line in light petrol.
def dark(deck, small: str, big: str, accent: str = "", notes: str = "", lines=(), dots=0):
    page = deck.slides.add_slide(deck.slide_layouts[6])
    shape(page, MSO_SHAPE.RECTANGLE, (0, 0, 13.333, 7.5), NAVY, "!!night")
    crowd = constellation(page, dots)
    top = say(page, small, (0.6, 2.1, 12.13, 0.6), 22, D.WHITE, center=True)
    head = page.shapes.add_textbox(*(Inches(x) for x in (0.6, 2.8, 12.13, 1.9)))
    for i, (words, colour) in enumerate(((big, D.WHITE), (accent, LIGHT))):
        if words:
            para = head.text_frame.paragraphs[0] if i == 0 else head.text_frame.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            D.run(para, words, 50, colour, bold=True)
    under = say(page, list(lines), (0.6, 4.85, 12.13, 1.8), 17, LIGHT, center=True)
    begin = 25 * len(crowd)
    steps = stagger(crowd, "zoom", 0, 25) + [(top, "fade", begin), (head, "zoom", begin + 400)]
    finish(page, notes, steps + ([(under, "rise", begin + 1300)] if lines else []))
    return page


# Sixty AI investor dots, each named for its agent so Morph flies it on.
def constellation(page, count: int) -> list:
    rng, dots = np.random.default_rng(7), []
    for i in range(count):
        band, k = divmod(i, 30)
        size = float(rng.uniform(0.14, 0.3))
        x = 0.45 + k * 0.42 + float(rng.uniform(-0.12, 0.12))
        y = (0.45 if band == 0 else 6.45) + float(rng.uniform(-0.3, 0.3))
        dots.append(agent(page, i, (x, y), size))
    return [dots[i] for i in rng.permutation(len(dots))]


# An investor dot in its persona's colour, or a copycat's recipe colour.
def agent(page, i: int, at, size: float):
    colour = PERSONAS[i // 5] if i < 30 else RECIPES[(i - 30) // 10]
    return shape(
        page, MSO_SHAPE.OVAL, (*at, size, size), RGBColor.from_string(colour), f"!!agent{i}"
    )


# The title slide: our line, team and schools over sixty investors.
def title(deck, f: dict, m: dict) -> None:
    lines = [
        "Market neutral. Hedged by machine learning. Watched by an AI crowd.",
        MEMBERS,
        SCHOOLS,
    ]
    meeting = f"{TEAM}  ·  McGill-FIAM Asset Management Hackathon 2026 Finals"
    dark(deck, meeting, "THE NEXT CROWD IS", "MADE OF MACHINES", notes("title", f), lines, dots=60)


# The hook: the brief's AI crowd warning, then the crowd we built.
def hook(deck, f: dict, m: dict) -> None:
    page = D.slide(deck, "Your Brief Warned About an AI Crowd", "", "The question")
    words = (
        "\u201cAgent-generated strategies tend to crowd into the same concentrated, high-valuation "
        "names \u2026 a recipe for exactly the kind of correlated unwind the industry last lived "
        "through in the August 2007 quant meltdown.\u201d"
    )
    quote = say(page, [words, "The organizers' 2026 brief"], (0.6, 1.75, 11.8, 2), 21)
    built = say(page, "So we built that crowd.", (0.6, 3.85, 12.13, 0.8), 40, D.ACCENT, True)
    # The title's sixty investors Morph in and settle beside the quote.
    for i in range(60):
        agent(page, i, (8.55 + (i % 15) * 0.28, 3.72 + (i // 15) * 0.28), 0.2)
    row = cards(
        page,
        [
            ("60", "AI investors we built"),
            (f["decisions_total"], "Decisions, every one on file"),
            ("0", "Bytes sent off our machine"),
        ],
        5.25,
    )
    steps = [(quote, "fade", 0), (built, "rise", 2600)] + stagger(row, "zoom", 4000, 300)
    finish(page, notes("hook", f), steps)


# Our answer in three pillars: theory, machine learning, the AI crowd.
def ideas(deck, f: dict, m: dict) -> None:
    page = D.slide(deck, "One Book, Three Ideas", "", "Our answer")
    steps = pillars(
        page,
        [
            (
                "Theory picks the stocks",
                "Twelve published signals. Nothing fitted to our test years.",
            ),
            (
                "Machine learning keeps it neutral",
                "It forecasts each stock's beta, so the hedge holds.",
            ),
            (
                "An AI crowd shows who we trade against",
                "Sixty AI investors we built and watch.",
            ),
        ],
    )
    finish(page, notes("ideas", f), steps)


# A stock earning its place, one real month's film (fiam.report.clips).
def journey(deck, f: dict, m: dict) -> None:
    movie, poster = CLIPS / "journey.mp4", CLIPS / "journey.png"
    told = load(CLIPS / "journey.json")
    if not (movie.exists() and poster.exists() and told):
        return journey_still(deck, f, m)
    page = deck.slides.add_slide(deck.slide_layouts[6])
    shape(page, MSO_SHAPE.RECTANGLE, (0, 0, 13.333, 7.5), NAVY, "!!night")
    clip = page.shapes.add_movie(
        str(movie), 0, 0, D.WIDE, D.TALL, poster_frame_image=str(poster), mime_type="video/mp4"
    )
    told = {k: f"{v:,}" if isinstance(v, int) else v for k, v in told.items()}
    finish(page, notes("journey", f, **told), film=clip)


# The same steps as circles, for a deck built without the film.
def journey_still(deck, f: dict, m: dict) -> None:
    page = D.slide(
        deck, "How a Stock Earns Its Place", "Every month, in this order", "How it works"
    )
    stages = [
        ("3,000", "US stocks every month"),
        ("12", "Published signals score each one"),
        ("1", "Score, best to worst"),
        ("0", "Net bet on the market or any sector"),
        (
            f"~{round(m['compliance']['positions_average'], -1):.0f}",
            "Names, the best long and the worst short",
        ),
    ]
    steps = []
    for i, (big, small) in enumerate(stages):
        x, ends = 0.75 + 2.5 * i, i in (0, 4)
        group = page.shapes.add_group_shape()
        shape(group, MSO_SHAPE.OVAL, (x, 2.0, 1.9, 1.9), D.ACCENT if ends else D.TINT)
        say(group, big, (x, 2.55, 1.9, 0.8), 34, D.WHITE if ends else D.ACCENT, True, True)
        say(group, small, (x - 0.15, 4.1, 2.2, 1.0), 15, D.INK, center=True)
        steps.append((group, "zoom", 700 * i))
        if i:
            arrow = shape(page, MSO_SHAPE.RIGHT_ARROW, (x - 0.5, 2.78, 0.42, 0.34), D.MUTED)
            steps.append((arrow, "wipe", 700 * i - 300))
    words = "Every input is dated before the decision. Every trading cost is charged."
    line = say(page, words, (0.6, 5.55, 12.13, 0.6), 20, D.ACCENT, True, True)
    finish(page, "", sorted(steps, key=lambda s: s[2]) + [(line, "rise", 3700)])


# Where machine learning won on beta and lost on volatility, pre-2021.
def ml(deck, f: dict, m: dict) -> None:
    title_, subtitle = (
        "Machine Learning Only Where It Earns Its Place",
        "Tested on the years before 2021",
    )
    page = D.slide(deck, title_, subtitle, "How it works")
    pic = chart(page, "ml_scoreboard", (0.6, 1.6, 12.13, 4.3))
    words = "It won on beta. It lost on volatility, so the simple rule runs there."
    line = say(page, words, (0.6, 6.05, 12.13, 0.7), 17, D.MUTED, center=True)
    finish(page, notes("ml", f), [(pic, "wipe", 0), (line, "fade", 1300)])


# The six AI colleagues, as a film when there is one, else as tiles.
def desk(deck, f: dict, m: dict) -> None:
    page = D.slide(deck, "Six AI Colleagues. None of Them Trades.", "", "How it works")
    movie, poster = CLIPS / "desk.mp4", CLIPS / "desk.png"
    if movie.exists() and poster.exists():
        box = (Inches(x) for x in (0.6, 1.45, 12.1, 4.6))
        clip = page.shapes.add_movie(
            str(movie), *box, poster_frame_image=str(poster), mime_type="video/mp4"
        )
        words = "Open models on our own machine. Every number checked."
        line = say(page, words, (0.6, 6.25, 12.13, 0.5), 17, D.ACCENT, True, True)
        finish(page, notes("desk", f), [(line, "rise", 3400)], film=clip)
        return
    at = (4.92, 1.7, 3.5, 0.85)
    boss = tile(page, at, "The portfolio manager", fill=D.ACCENT, ink=D.WHITE, size=18)
    roles = [
        ("Rules reader", "Quotes every limit"),
        ("Triage", "Sorts company filings"),
        ("Analyst", "Reads a stock's month"),
        ("Devil's advocate", "Attacks the book"),
        ("Memo writer", "Checks every number"),
        ("Explainer", "Says why we held a name"),
    ]
    row = [
        tile(page, (0.6 + 4.1 * (i % 3), 2.95 + 1.5 * (i // 3), 3.85, 1.3), *role, size=19)
        for i, role in enumerate(roles)
    ]
    words = "Open models on our own machine. Every number checked."
    line = say(page, words, (0.6, 6.15, 12.13, 0.5), 17, D.ACCENT, True, True)
    steps = [(boss, "zoom", 0)] + stagger(row, "zoom", 500, 180) + [(line, "rise", 1800)]
    finish(page, notes("desk", f), steps)


# The navy turn to results: did it work, January 2021 to August 2026.
def results(deck, f: dict, m: dict) -> None:
    small = "January 2021 to August 2026. Every cost charged."
    dark(deck, small, "So, did it", "WORK?", notes("results", f))


# Growth of $1 vs the hurdle and S&P 500, with the IR and net return.
def growth(deck, f: dict, m: dict) -> None:
    perf = m["performance"]
    subtitle = "Growth of $1 against cash plus 4% and the S&P 500, 2021 to August 2026"
    page = D.slide(deck, "Yes. It Beat the Hurdle After Every Cost", subtitle, "Results")
    pic, film = moving(page, "growth", (0.6, 1.65, 8.3, 5.2), "cumulative_return", book=True)
    items = [
        (f"{perf['information_ratio']:.2f}", "IR above cash plus 4%"),
        (f"{perf['return_net_annual']:.1%}", "Net return a year after every cost"),
    ]
    row = column(page, items, 9.2, 1.8, 3.53, 2.15, 2.45, 54)
    steps = [] if film else [(pic, "wipe", 0)]
    steps += stagger(row, "zoom", 3700 if film else 1200, 400)
    finish(page, notes("growth", f), steps, film=film)


# Rolling S&P 500 beta, called zero only within two standard errors.
def neutral(deck, f: dict, m: dict) -> None:
    perf = m["performance"]
    beta, error = perf["beta_realized"], perf["beta_stderr"]
    zero = abs(beta) < 2 * error
    heading = "And It Never Bet on the Market" if zero else "Its Bet on the Market, Measured"
    subtitle = "Rolling 12-month beta to the S&P 500, 2021 to August 2026"
    page = D.slide(deck, heading, subtitle, "Results")
    pic, film = moving(page, "beta", (0.6, 1.65, 8.3, 5.2), "rolling_beta", book=True)
    label = "Beta to the S&P 500" + (", statistically zero" if zero else "")
    hero = card(page, (9.2, 1.8, 3.53, 2.15), f"{beta:+.2f}", label, 54)
    rho = perf["correlation_with_market"]
    moves = card(page, (9.2, 4.2, 3.53, 2.15), f"{rho:+.2f}", "Correlation with the S&P 500", 54)
    start = 3600 if film else 1200
    steps = ([] if film else [(pic, "wipe", 0)]) + [
        (hero, "zoom", start),
        (moves, "zoom", start + 500),
    ]
    words = notes("neutral", f, signed_beta=f"{beta:+.2f}", rho=f"{rho:+.2f}")
    finish(page, words, steps, film=film)


# Dollars and IR by fund size, where v1 peaks and stops, final at $5B.
def capacity(deck, f: dict, m: dict) -> None:
    subtitle = "Dollars earned and IR by fund size, after all costs"
    page = D.slide(deck, "How Much Money Can It Run?", subtitle, section("Capacity", True))
    pic, film = moving(page, "capacity", (0.6, 1.55, 12.1, 3.6), "capacity")
    items = [
        (f["v1_peak"], "Our first book peaks"),
        (f["v1_breakeven"], "Our first book stops paying"),
        ("$5B", f"The final book still earns an IR of {f['fv3_ir_5e9']}"),
    ]
    steps = [] if film else [(pic, "wipe", 0)]
    steps += stagger(cards(page, items, 5.25), "zoom", 3600 if film else 1300, 300)
    finish(page, notes("capacity", f), steps, film=film)


# Final v3 (v2_amendment_3.md) against v1, v2 and v2-pass by fund size.
def final(deck, f: dict, m: dict) -> None:
    every = f["fd_best_sharpe"] == f["fd_sizes"]
    heading = "One Book for Every Fund Size, $10M to $5B" if every else "The Book We Present"
    subtitle = "Against our first book and two alternatives at every fund size, after all costs"
    page = D.slide(deck, heading, subtitle, section("The final book", True))
    pic = chart(page, "final_book", (0.6, 1.5, 12.13, 3.25))
    se = f["fv3_beta_se"]
    zero = se != "-" and abs(float(f["fv3_beta_1e7"])) < 2 * float(se)
    best = (
        "Sharpe, best of the four at every size"
        if every
        else f"Sharpe, best of the four at {f['fd_best_sharpe']} of {f['fd_sizes']} sizes"
    )
    items = [
        (f["fv3_ir_1e7"], "IR at $10M"),
        (f["fv3_ir_5e9"], "IR at $5B"),
        (f["fv3_beta_1e7"], "Beta" + (", statistically zero" if zero else "")),
        (f["fv3_sharpe_1e7"], best),
    ]
    steps = [(pic, "wipe", 0)] + stagger(cards(page, items, 5.05), "zoom", 1300, 300)
    words = (
        f"The highest deflated Sharpe of the four at {f['fd_best_dsr']} of {f['fd_sizes']} sizes, "
        f"{f['dsr_fv3']} at $10M after {f['trials']} trials. Every rule was written down before "
        f"it ran, and the ones that failed their own tests are out, {f['fv3_dropped']}."
    )
    if MODE == "print":
        steps.append((say(page, words, (0.6, 6.55, 12.13, 0.4), 12, D.MUTED), "fade", 2600))
    finish(page, notes("final", f), steps)


# Where it struggles: weak IR years beside the crisis replays.
def losses(deck, f: dict, m: dict) -> None:
    subtitle = "After all costs, with 2016 to 2020 as the pre-test"
    page = D.slide(deck, "Where It Struggles, and Why", subtitle, section("Risk", True))
    items = [
        (f["fv3_check_ir"], "IR 2016 to 2020, a short tilt in a rising market"),
        (
            f"{f['ir_2021_22']} then {f['ir_2023_26']}",
            "IR 2021 to 2022, then 2023 to 2026 through a junk rally",
        ),
    ]
    row = column(page, items, 0.6, 1.75, 5.6, 2.0, 2.3, 44)
    pic = chart(page, "crises", (6.5, 1.55, 6.23, 5.4))
    crises = load(V2 / "stress.json").get("final v3", {}).get("crises", {})
    lost = [c["loss"] for c in crises.values()]
    replay = dict(
        crises=len(lost),
        crises_up=sum(x > 0 for x in lost),
        crisis_worst=f"{-min(lost, default=0):.1%}",
    )
    said = notes("losses", f, **replay)
    finish(page, said, stagger(row, "rise", 0, 400) + [(pic, "wipe", 1200)])


# The navy turn to the edge tests: is the edge real?
def edge(deck, f: dict, m: dict) -> None:
    dark(deck, "Every backtest looks good.", "Is the edge", "REAL?", notes("edge", f))


# Final v3 among 1,000 random large-cap books (final_exhibits.md 1.1).
def monkeys(deck, f: dict, m: dict) -> None:
    subtitle = "Long-short books of the 500 largest stocks, with our size, holdings and costs"
    page = D.slide(
        deck, "1,000 Random Books. None Came Close.", subtitle, section("The edge", True)
    )
    pic, film = moving(page, "null", (0.6, 1.6, 8.3, 4.9), "edge_null")
    items = [
        (f["null_beaten"], "Random books beaten"),
        (f["null_best"], "IR of the best random book"),
        (f["fv3_ir_1e7"], "IR of the final book"),
    ]
    row = column(page, items, 9.2, 1.75, 3.53, 1.45, 1.65, 34)
    steps = [] if film else [(pic, "wipe", 0)]
    steps += stagger(row, "zoom", 3500 if film else 1200, 350)
    finish(page, notes("monkeys", f), steps, film=film)


# Every benchmark on one axis, and the live outlook (final_exhibits.md).
def ladder(deck, f: dict, m: dict) -> None:
    subtitle = "IR at $10M after all costs, 2021 to August 2026"
    page = D.slide(deck, "Every Benchmark We Could Think Of", subtitle, section("The edge", True))
    pic, film = moving(page, "ladder", (0.6, 1.55, 8.4, 4.4), "edge_ladder")
    items = [
        (f"{f['live_low']} to {f['live_high']}", "Expected live IR after haircuts"),
        (f["dsr_fv3"], "Chance it is skill, after every trial"),
        (f["trials"], "Trials counted, every one"),
    ]
    row = column(page, items, 9.3, 1.75, 3.43, 1.45, 1.65, 30)
    steps = [] if film else [(pic, "wipe", 0)]
    steps += stagger(row, "zoom", 3500 if film else 1300, 350)
    finish(page, notes("ladder", f), steps, film=film)


# The printed edge page: the random books and every benchmark.
def edge_page(deck, f: dict, m: dict) -> None:
    subtitle = "The final book at $10M, 2021 to August 2026, after all costs"
    page = D.slide(deck, "Is the Edge Real? Random Books and Every Benchmark", subtitle, "Appendix")
    chart(page, "edge_null", (0.6, 1.55, 6.0, 3.0))
    chart(page, "edge_ladder", (6.75, 1.55, 5.98, 3.0))
    words = (
        f"The final book beats {f['null_beaten']} random large-cap books, whose average is "
        f"{f['null_mean']} and best {f['null_best']}. Our rule on the 500 largest alone makes "
        f"{f['large_10m']} at $10M and {f['large_1bn']} at $1B. Our favourite names held still "
        f"with hindsight make {f['fav_10m']}, and equal weights on our score "
        f"{f['naive_model_10m']}. "
        f"Expected live after published haircuts, {f['live_low']} to {f['live_high']}."
    )
    D.text(page, [words], D.LEFT, Inches(4.85), D.WIDTH, Inches(1.8), 12)
    finish(page)


# The printed crowd page: how the AI crowd herds and copycats find us.
def crowd_page(deck, f: dict, m: dict) -> None:
    title_ = f"30 AI Investors Vote Like {f['voters']}, and Copycats Land on Our Stocks"
    page = D.slide(deck, title_, CLEAN, "Appendix")
    chart(page, "crowd_network", (0.6, 1.55, 4.0, 3.6))
    chart(page, "copycats", (4.85, 1.55, 7.88, 3.6))
    items = [
        (f["voters"], "Independent voters among 30 retail AIs"),
        (f["copy_longs"], "How often copycats given our recipe buy our longs, against chance"),
        (f["fv3_crowd_beta"], "The final book's loading on the AI crowd's favourites"),
    ]
    cards(page, items, 5.5)
    finish(page)


# Final v3's risk, name by name (final_exhibits.md).
def risk_slide(deck, f: dict, m: dict) -> None:
    subtitle = "Every month, name by name, from our risk model"
    page = D.slide(deck, "Where the Risk Sits", subtitle, section("Risk", True))
    pic = chart(page, "final_risk", (0.6, 1.5, 12.13, 3.75))
    items = [
        (f["risk_vol"], "Forecast volatility a year"),
        (f["risk_factor_share"], "Risk from styles"),
        (f["risk_top10"], "Risk in the top 10 names"),
        (f["risk_var95"], "Monthly VaR at 95%"),
    ]
    row = cards(page, items, 5.45, height=1.2, size=30)
    finish(page, notes("risk", f), [(pic, "wipe", 0)] + stagger(row, "zoom", 1300, 300))


# Final v3's bets by month, told by the explainer, every number checked.
def exposures_slide(deck, f: dict, m: dict) -> None:
    subtitle = "Exposure to five styles, every month"
    page = D.slide(deck, "What It Bets On, Month by Month", subtitle, section("Risk", True))
    pic = chart(page, "final_exposures", (0.6, 1.55, 12.13, 3.6))
    first = f["factor_story"].split(". ")[0].rstrip(".") + "."
    quote = [f"\u201c{first}\u201d", "Our explainer agent, every number checked"]
    story = say(page, quote, (0.6, 5.45, 12.13, 1.3), 17, D.ACCENT, center=True)
    finish(page, notes("exposures", f), [(pic, "wipe", 0), (story, "fade", 1400)])


# The printed risk page: risk table, exposures, what the optimizer adds.
def risk_page(deck, f: dict, m: dict) -> None:
    subtitle = "The final book, 2021 to August 2026, after all costs, risk from our factor model"
    heading = "Where the Risk Sits, What It Bets On, What the Optimizer Adds"
    page = D.slide(deck, heading, subtitle, "Appendix")
    chart(page, "final_risk", (0.6, 1.5, 12.13, 2.75))
    chart(page, "final_exposures", (0.6, 4.35, 7.2, 2.4))
    chart(page, "naive", (8.05, 4.3, 4.68, 2.45))
    finish(page)


# The navy turn to the crowd: who is on the other side of our trades.
def turn(deck, f: dict, m: dict) -> None:
    small = "Who is on the other side of our trades?"
    dark(deck, small, "Meet", "THE AI CROWD", notes("turn", f), dots=60)


# The crowd by kind: retail personas in rows, copycats by their recipe.
def meet(deck, f: dict, m: dict) -> None:
    from fiam.agents import crowd_lab

    page = D.slide(deck, "Sixty AI Investors, Each With Its Own View", "", "The AI crowd")
    labels = [say(page, "30 retail investors", (0.6, 1.5, 5.6, 0.4), 18, D.ACCENT, True)]
    for r, persona in enumerate(crowd_lab.PERSONAS):
        labels.append(say(page, persona.capitalize(), (0.6, 2.02 + 0.72 * r, 1.9, 0.4), 15))
        for c in range(5):
            agent(page, 5 * r + c, (2.55 + 0.72 * c, 1.95 + 0.72 * r), 0.5)
    labels.append(say(page, "30 copycat funds", (6.95, 1.5, 5.8, 0.4), 18, D.ACCENT, True))
    for r, recipe in enumerate(("Given our recipe", "Given the 2026 brief", "Given the 2025 case")):
        labels.append(say(page, recipe, (6.95, 1.98 + 1.32 * r, 5.8, 0.4), 15))
        for c in range(10):
            agent(page, 30 + 10 * r + c, (6.95 + 0.57 * c, 2.42 + 1.32 * r), 0.44)
    words = (
        "Each sees anonymous stock cards and picks eight. No names, no dates, and only months "
        "after its training ended."
    )
    line = say(page, words, (0.6, 6.25, 12.13, 0.5), 16, D.MUTED, center=True)
    finish(page, notes("meet", f), stagger(labels, "fade", 600, 120) + [(line, "rise", 1800)])


# Retail investors fly into the groups their votes form, biggest first.
def swarm(deck, f: dict, m: dict) -> None:
    net = load(V2 / "crowd_insights.json").get("network", {})
    groups = dict(zip(net.get("agents", []), net.get("groups", [])))
    counts = {g: list(groups.values()).count(g) for g in set(groups.values())}
    sizes = sorted(counts.items(), key=lambda kv: -kv[1])
    page = D.slide(deck, "Thirty Voices, One Big Herd", "", "The AI crowd")
    if not groups:
        say(page, "crowd network not measured yet", (0.6, 3, 12, 0.5), 14, D.MUTED)
        finish(page)
        return
    labels, x = [], 1.2
    for name, count in sizes:
        radius = 0.3 * np.sqrt(count) + 0.3
        cx, cy = x + radius, 3.3
        members = [a for a in sorted(groups) if groups[a] == name]
        for k, a in enumerate(members):
            r, spin = 0.3 * np.sqrt(k), k * 2.39996
            agent(page, int(a), (cx + r * np.cos(spin) - 0.25, cy + r * np.sin(spin) - 0.25), 0.5)
        words = f"{count} agent" + ("s" if count > 1 else "")
        labels.append(
            say(page, words, (cx - 1, cy + radius + 0.25, 2, 0.4), 16, D.MUTED, True, True)
        )
        x += 2 * radius + 0.9
    biggest = sizes[0][1]
    words = (
        f"{biggest} of 30 vote together. Six personas and five kinds of data, yet they vote like "
        f"{f['voters']} independent investors. What they see matters more than who they are "
        "told to be."
    )
    line = say(page, words, (0.6, 6.05, 12.13, 0.8), 17, D.ACCENT, True, True)
    steps = stagger(labels, "fade", 1200, 150) + [(line, "rise", 2000)]
    finish(page, notes("swarm", f, biggest=biggest), steps)


# The AI Crowd Lab film full screen, the talk's cut if there is one.
def film(deck, f: dict, m: dict) -> None:
    page = deck.slides.add_slide(deck.slide_layouts[6])
    shape(page, MSO_SHAPE.RECTANGLE, (0, 0, 13.333, 7.5), NAVY, "!!night")
    movie = FIGURES / "ai_crowd_short.mp4"
    movie = movie if movie.exists() else FIGURES / "ai_crowd.mp4"
    poster = FIGURES / "ai_crowd_poster.png"
    if not (movie.exists() and poster.exists()):
        say(page, "The AI Crowd Lab film", (0.6, 3.2, 12.13, 0.8), 34, D.WHITE, True, True)
        finish(page, notes("film", f))
        return
    clip = page.shapes.add_movie(
        str(movie), 0, 0, D.WIDE, D.TALL, poster_frame_image=str(poster), mime_type="video/mp4"
    )
    finish(page, notes("film", f), film=clip)


# How often copycats buy our longs, given the brief or our own recipe.
def copycats(deck, f: dict, m: dict) -> None:
    from fiam.report.figures import copycat_test

    test = copycat_test()
    brief = test.get("brief_2026", {}).get("buys_our_longs")
    recipe = test.get("v1", {}).get("buys_our_longs")
    subtitle = (
        "How often AI copycats buy our longs against chance, after the model's training ended"
    )
    heading = "Hand Any AI a Quant Brief and It Lands on Our Stocks"
    page = D.slide(deck, heading, subtitle, section("AI Crowd Lab", True))
    pic, film = moving(page, "copycats", (0.6, 1.6, 7.6, 5.2), "copycats")
    items = [
        (f"{brief:.1f}x" if brief else "-", "Our longs bought, given only the organizers' brief"),
        (f"{recipe:.1f}x" if recipe else "-", "Given our own recipe"),
    ]
    row = column(page, items, 8.6, 1.8, 4.13, 2.1, 2.4, 50)
    steps = [] if film else [(pic, "wipe", 0)]
    steps += stagger(row, "zoom", 3300 if film else 1200, 400)
    finish(page, notes("copycats", f), steps, film=film)


# Our crowding beta and our loss if a $50B AI crowd sold at once.
def trades(deck, f: dict, m: dict) -> None:
    subtitle = "The AI crowd's real votes over 19 months"
    page = D.slide(deck, "We Are Already in Their Trades", subtitle, section("AI Crowd Lab", True))
    items = [
        (f["fv3_crowd_beta"], "Crowding beta to the AI crowd's favourites"),
        (f["fv3_sell_50"], "Our loss if a $50B AI crowd sold at once"),
    ]
    row = cards(page, items, 1.75, height=1.5, size=44)
    pic, film = moving(page, "overlap", (0.6, 3.5, 12.1, 3.4), "crowd_signature")
    steps = [] if film else [(pic, "wipe", 1100)]
    steps += stagger(row, "zoom", 3700 if film else 0, 400)
    finish(page, notes("trades", f), steps, film=film)


# The copycat's recipe word for word, without its list of signals.
def told_short(recipe: str) -> str:
    head, _, rest = recipe.partition(":")
    last = rest.split(". ", 1)[1] if ". " in rest else ""
    return f"“{head} … {last}”" if last else f"“{recipe}”"


# A measure's name fit for a column: before any colon, cut at a word.
def measure(name: str) -> str:
    name = name.split(":")[0].strip()
    return name if len(name) <= 26 else name[:26].rsplit(" ", 1)[0]


# A saved decision rebuilt: what the copycat was told, saw and answered.
def receipt(deck, f: dict, m: dict) -> None:
    r = load(FIGURES / "receipt.json")
    same = r.get("hash_match") and r.get("same_stocks")
    month = f"{pd.Timestamp(r['month']):%B %Y}" if r.get("month") else ""
    subtitle = f"One AI copycat given our recipe in {month}, rebuilt from the data today. " + (
        "Same 40 stocks, same prompt hash" if same else "The rebuild does not match"
    )
    page = D.slide(
        deck, "Every Decision Is on File. Here Is One.", subtitle, section("Trust", True)
    )
    if not r:
        finish(page)
        return
    told = [
        say(page, "What it was told", (0.6, 1.6, 5.6, 0.3), 13, D.ACCENT, True),
        say(page, told_short(r["recipe"]), (0.6, 1.95, 5.6, 1.5), 11, D.MUTED),
    ]
    head = "| Code | " + " | ".join(measure(c) for c in r["columns"][:3]) + " |"
    body = "\n".join("| " + " | ".join(str(v) for v in row[:4]) + " |" for row in r["rows"])
    stocks, measures = r["size"]
    note = f"{stocks} stocks and {measures} measures, percentile ranks only. No names, no dates."
    seen = f"{head}\n|---|---|---|---|\n{body}\n\n{note}"
    saw = [say(page, "What it saw", (0.6, 3.5, 5.6, 0.3), 13, D.ACCENT, True)]
    saw += added(page, D.table, seen, Inches(0.6), Inches(3.85), Inches(5.6), 9, 0.24)
    rows = ["| Code | It said | Stock | Our book |", "|---|---|---|---|"]
    rows += [
        f"| {p['code']} | {p['side'].capitalize()} | {p['ticker']}, {p['name']} | "
        f"{p['ours'].capitalize()} |"
        for p in r["picks"]
    ]
    heading = "What it answered, and who they were"
    said = [say(page, heading, (6.6, 1.6, 6.1, 0.3), 13, D.ACCENT, True)]
    said += added(
        page, D.table, "\n".join(rows), Inches(6.6), Inches(1.95), Inches(6.13), 8.5, 0.22
    )
    hits = sum((p["side"], p["ours"]) in (("buy", "long"), ("short", "short")) for p in r["picks"])
    words = (
        f"{hits} of its {len(r['picks'])} calls match our book. Same card, same answer "
        f"{f['rerun']} of the time. Nudge the numbers and only {f['moved']} of picks survive."
    )
    proof = say(page, words, (0.6, 6.55, 12.13, 0.4), 13, D.ACCENT, True)
    steps = stagger(told, "fade", 0, 300) + stagger(saw, "rise", 1000, 200)
    steps += stagger(said, "wipe", 2200, 300) + [(proof, "zoom", 3600)]
    finish(page, notes("receipt", f), steps)


# Today's finals-rules book from real data only (finals_global.md).
def today(deck, f: dict, m: dict) -> None:
    runs = sorted((V2 / "finals").glob("*.json"), key=lambda p: p.stat().st_mtime)
    real = [load(p) for p in runs]
    real = [r for r in real if r.get("final v3", {}).get("data", "world_fake") != "world_fake"]
    if not real:
        return
    books = real[-1]
    best, rules = books["final v3"], books["final v3"]["mandate"]
    countries = ", ".join(rules["countries"])
    subtitle = (
        f"{rules['gross']:.0%} gross, {rules['net']:+.0%} net, {rules['names_min']} to "
        f"{rules['names_max']} names, {countries}, decisions {best['window'][0][:7]} to "
        f"{best['window'][1][:7]}, after all costs"
    )
    page = D.slide(deck, "Under Today's Rules", subtitle, "The finals")
    beta = best.get("beta_world", best["beta_realized"])
    items = [
        (f"{best['information_ratio']:.2f}", "IR of the final book"),
        (f"{beta:+.2f}", "Beta to the market"),
        (f"{best['positions']:.0f}", "Names on average"),
        (str(best["months_breaching"] + best["months_empty"]), "Months breaching a rule"),
    ]
    row = cards(page, items, 1.75, height=1.7, size=44)
    world = "country_share_max" in best
    head = "| Book | IR | Costs a Year | Names | Months Breaching |"
    rows = [head + (" Largest Country |" if world else ""), "|---|" + "---:|" * (4 + world)]
    shown = {"final v3": "Final book", "v1 method": "First book's method", "v2": "Built for size"}
    for label in sorted(books, key=lambda k: k != "final v3"):
        b = books[label]
        rows.append(
            f"| {shown.get(label, label)} | {b['information_ratio']:.2f} | "
            f"{b['cost_annual']:.2%} | {b['positions']:.0f} | "
            f"{b['months_breaching'] + b['months_empty']} |"
            + (f" {b['country_share_max']:.0%} |" if world else "")
        )
    D.table(page, "\n".join(rows), Inches(0.6), Inches(3.9), Inches(12.13), 12, 0.36)
    finish(page, notes("today", f), stagger(row, "zoom", 0, 300))


# What we would build next: a growing book, a crowd monitor, fast rules.
def launch(deck, f: dict, m: dict) -> None:
    page = D.slide(deck, "What We Would Build Next", "", "The product")
    steps = pillars(
        page,
        [
            (
                "A book that grows with you",
                "One book that beats the hurdle from $10M to $5B."
                if f["v3_size"] == "pass"
                else f"Our first book below {f['crossover']}; the second, built for size, "
                "above it.",
            ),
            (
                "A monitor for the AI crowd",
                "Our exposure to the AI crowd, and its sell-off cost.",
            ),
            (
                "Rules to portfolio in minutes",
                "New rules to a compliant book in minutes.",
            ),
        ],
    )
    finish(page, notes("launch", f), steps)


# The closing navy slide: what we would do with real money, and thanks.
def closing(deck, f: dict, m: dict) -> None:
    lines = [
        "One book that beats the hurdle from $10M to $5B."
        if f["v3_size"] == "pass"
        else f"Our first book below {f['crossover']}, the second above it.",
        "Watch the AI crowd like any other risk.",
        "Every test written down before it ran.",
        f"Thank you from {TEAM}",
    ]
    meeting = "What we would do with real money"
    dark(deck, meeting, "Run it, and", "WATCH THE CROWD", notes("closing", f), lines, dots=60)


# Page 1 of the brief.
def summary(deck, f: dict, m: dict) -> None:
    perf = m["performance"]
    side = R.load("exhibits/pack.json")["completeness"]["comparison"]
    heading = "A Market-Neutral Book, Hedged by Machine Learning, Watched by an AI Crowd"
    subtitle = f"{D.PERIOD}. Benchmark: 3-month T-bill + 4% a year. S&P 500 for context"
    page = D.slide(deck, heading, subtitle, section("Executive summary"))
    items = [
        (f"{perf['information_ratio']:.2f}", "information ratio vs T-bill + 4%"),
        (f"{perf['return_net_annual']:.1%}", "net return a year, after all costs"),
        (
            f"{perf['beta_realized']:+.2f}",
            f"beta to the S&P 500 (standard error {perf['beta_stderr']:.2f})",
        ),
        (f"{perf['max_drawdown_total']:.1%}", "worst drawdown, monthly"),
    ]
    cards(page, items, 1.6)
    text = [
        (
            "Strategy.",
            "Long the strongest US stocks and short the weakest on twelve published "
            "signals, dollar, beta and sector neutral.",
        ),
        (
            "Machine learning.",
            "LightGBM forecasts each stock's beta so the hedge holds. Where "
            "machine learning lost to a simple model, the simple model runs.",
        ),
        (
            "Agents.",
            "Six local AI roles read rules, check numbers and explain; none sets a "
            "position. Sixty AI investors show who we trade against, and the 30 retail ones "
            f"vote like {f['voters']} independent investors.",
        ),
        (
            "Integrity.",
            "Every input dated before the decision, every cost charged, every AI "
            "answer saved and replayable.",
        ),
    ]
    D.text(page, text, D.LEFT, Inches(3.15), D.HALF, Inches(3.7), 13, gap=9)
    glance = D.glance(side, perf, m["costs"], m["compliance"])
    D.table(page, glance, D.RIGHT, Inches(3.15), D.HALF, 11.5, 0.34)
    finish(page)


# Brief page 2: legs, enforced neutrality, signals, top names, growth.
def strategy(deck, f: dict, m: dict) -> None:
    comp = m["compliance"]
    side = R.load("exhibits/pack.json")["completeness"]["ranges"]["gross_exposure"]["mean"] / 2
    heading = "Long the Best, Short the Worst, Across Twelve Published Signals"
    stocks = f"{comp['positions_average']:.0f} stocks on average"
    subtitle = f"Market neutral by design, {stocks}, rebalanced monthly"
    page = D.slide(deck, heading, subtitle, section("Strategy"))
    signals = list(D.NAMES.values())
    sized = (
        "Positions capped at 1.5% of NAV and 40% of a day's trading, trading costs at 0.165% of "
        "NAV a month, risk aimed at 10% a year, gross cut to 100% after a 5% fall."
    )
    text = [
        (
            "Score.",
            "Each month every stock is ranked on twelve published signals, equally "
            "weighted: " + ", ".join(signals[:-1]) + f" and {signals[-1]}.",
        ),
        (
            "Neutral.",
            f"{side:.0%} long and {side:.0%} short, so dollar neutral (net band ±50%) at "
            f"{2 * side:.0%} gross (limit 200%). Beta held within ±0.005 on LightGBM's "
            "forecast, each sector within ±1%.",
        ),
        ("Sized.", sized),
    ]
    D.text(page, text, D.LEFT, Inches(1.7), D.HALF, Inches(2.4), 12.5, gap=7)
    D.chart(page, "cumulative_return", D.LEFT, Inches(4.15), D.HALF, Inches(2.8))
    top = D.renamed(R.top(m)).replace("| weight |", "| average weight |")
    named = R.load("exhibits/labels.json")
    rows = named.get("holding_rows", 0)
    note = "Average weight over every month, zero when not held."
    if rows:
        sure = round(rows * named.get("confirmed_share", 0))
        note += f" {sure:,} of {rows:,} holding rows confirmed by CRSP, an 8-K, EDGAR or Compustat."
    D.table(page, top + "\n\n" + note, D.RIGHT, Inches(1.75), D.HALF, 9.5, 0.29)
    finish(page)


# Brief page 3: data, method, out-of-sample R² and each agent's setup.
def method(deck, f: dict, m: dict) -> None:
    with submitted():
        mine = R.load("metrics.json")
        subtitle = (
            "Every model is trained only on data from before the year it predicts, and refit each "
            "January"
        )
        page = D.slide(
            deck, "Every Model Learns Only From the Past", subtitle, section("Data and methodology")
        )
        D.boxes(page, D.BOOK_PATH, Inches(1.6), Inches(0.75), D.TINT)
        error = re.search(r"(\d\.\d\d) against (\d\.\d\d)", R.teachers(mine))
        cut = f", cutting the 60-month beta's error from {error[2]} to {error[1]}" if error else ""
        why = (
            "Return models follow the brief: train from 2015, validate on the two years before, "
            "test the next year. An 8-K is used only from the month after it is filed. LightGBM "
            "for beta because beta shifts with size, volatility and leverage in ways a line "
            f"misses{cut}."
        )
        below = D.table(page, R.r2(mine), D.LEFT, Inches(2.55), Inches(4.6), 9, 0.25)
        D.text(page, [why], D.LEFT, below + Inches(0.08), Inches(4.6), Inches(1.2), 9.5, D.MUTED)
        one, gemma = "None, one answer each", "Gemma 4 26B-A4B"
        extra = [
            (
                "AI Crowd Lab, 30 retail",
                gemma,
                "Cards of 40 anonymous stocks, percentiles only, after its training ended",
                one,
                f"They vote like {f['voters']} independent investors",
            ),
            (
                "AI copycats, 30 funds",
                gemma,
                "The same cards and one paragraph of a strategy",
                one,
                f"Given our recipe, buy our longs {f['copy_longs']} as often as chance",
            ),
            (
                "Rules reader",
                gemma,
                "A rules document, in pieces",
                "One answer a piece, every quote checked",
                f"{f['rules']} of last year's limits found, each quoted",
            ),
        ]
        # Deck 1's crowd row counts v1's longs; the finals deck says whose.
        crowd = D.architecture().replace("of our longs are", "of the first book's longs are")
        table = crowd + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in extra)
        below = D.table(page, table, Inches(5.45), Inches(2.55), Inches(7.28), 8, 0.23)
        firewall = (
            f" The AI Crowd Lab sees only months after Gemma's training ended; the same card gives "
            f"the same answer ({f['rerun']}) and moved numbers move the picks (only {f['moved']} "
            "survive), so it reads the card rather than remembering the stock."
        )
        top = below + Inches(0.05)
        D.text(page, [D.recall() + firewall], Inches(5.45), top, Inches(7.28), Inches(1.0), 9)
        finish(page)


# Brief page 4: returns vs the hurdle and market, risk, exposure, costs.
def returns_page(deck, f: dict, m: dict) -> None:
    perf = m["performance"]
    side = R.load("exhibits/pack.json")["completeness"]["comparison"]
    heading = (
        f"Information Ratio {perf['information_ratio']:.2f} After Costs, Alpha "
        f"{perf['alpha_annual']:.1%} a Year (t-stat {perf['alpha_tstat']:.2f})"
    )
    page = D.slide(deck, heading, D.PERIOD, section("Performance"))
    half = R.load("exhibits/pack.json")["completeness"]["ranges"]["gross_exposure"]["mean"] / 2
    sides = f"{half:.0%} long and {half:.0%} short on average"
    returns = D.returns(side, R).replace("80% long and 80% short", sides)
    D.table(page, returns, D.LEFT, Inches(1.75), D.HALF, 10, 0.28)
    D.table(page, D.risk(R, m), D.RIGHT, Inches(1.75), D.HALF, 10, 0.27)
    finish(page)


# Page 5: growth, drawdowns and every year against the hurdle.
def years_page(deck, f: dict, m: dict) -> None:
    perf = m["performance"]
    heading = (
        f"Ahead of the Hurdle in {D.hurdle_years(R, m)}, Worst Drawdown "
        f"{perf['max_drawdown_total']:.1%}"
    )
    page = D.slide(deck, heading, D.PERIOD, section("Performance"))
    D.chart(page, "cumulative_return", D.LEFT, Inches(1.75), D.HALF)
    D.chart(page, "drawdown", D.RIGHT, Inches(1.75), D.HALF)
    D.table(page, R.years(m), D.LEFT, Inches(5.0), D.WIDTH, 10.5, 0.27)
    finish(page)


# Page 6: beta month by month, and day by day with a lead and a lag.
def beta_page(deck, f: dict, m: dict) -> None:
    perf = m["performance"]
    dimson = load(V2 / "evidence.json").get("final_daily", {}).get("dimson_beta")
    daily = (
        f"Measured on daily returns with a lead and a lag (2021-2025), beta is {dimson:+.2f}"
        if dimson is not None
        else "Daily beta not measured for this book"
    )
    beta, error = perf["beta_realized"], perf["beta_stderr"]
    zero = ", Statistically Zero" if abs(beta) < 2 * error else ""
    heading = f"Monthly Beta {beta:+.2f} (Standard Error {error:.2f}){zero}"
    page = D.slide(
        deck, heading, f"Rolling 12-month windows, {D.PERIOD}. {daily}", section("Neutrality")
    )
    quarter, high, low, tall = Inches(5.6), Inches(1.6), Inches(4.35), Inches(2.65)
    D.chart(page, "rolling_beta", D.LEFT, high, quarter, tall)
    D.chart(page, "rolling_ir", D.RIGHT, high, quarter, tall)
    D.chart(page, "rolling_active_return", D.LEFT, low, quarter, tall)
    D.chart(page, "return_histogram", D.RIGHT, low, quarter, tall)
    finish(page)


# Page 7: the brief's limits as measured, every cost, and the two legs.
def costs_page(deck, f: dict, m: dict) -> None:
    cost = m["costs"]
    heading = (
        f"All Costs Modelled: {cost['cost_annual']:.1%} a Year, "
        f"{cost['cost_share_of_gross']:.0%} of Gross Return"
    )
    page = D.slide(deck, heading, D.PERIOD, section("Implementation"))
    limits = R.compliance(m).replace("| ±0.02 |", "| ±0.005 |")
    D.table(page, limits, D.LEFT, Inches(1.75), D.HALF, 10, 0.29)
    D.table(page, R.costs(m), D.RIGHT, Inches(1.75), D.HALF, 10, 0.29)
    D.table(page, R.legs(m), D.LEFT, Inches(4.8), D.WIDTH, 10, 0.29)
    finish(page)


# Pages 4 to 7 of the brief: deck 1's committee pack for final v3.
def committee(deck, f: dict, m: dict) -> None:
    for page in (returns_page, years_page, beta_page, costs_page):
        page(deck, f, m)


# Brief page 8: deck 1's discussion, plus what crowd and capacity add.
def discussion(deck, f: dict, m: dict) -> None:
    try:
        lines = D.discussion(R, m)
    except KeyError, IndexError, StopIteration, TypeError:
        perf = m["performance"]
        lines = [
            f"Expectations: IR {perf['information_ratio']:.2f} with beta "
            f"{perf['beta_realized']:+.2f} over the full period."
        ]
    # Deck 1's page numbers point at its own appendix, not this deck's.
    lines = [re.sub(r" \(pages? [\d, and-]+\)", "", line) for line in lines]
    lines = [line for line in lines if not line.startswith("Next:")]
    lines += [
        f"Crowding: {f['fv3_overlap_retail']} of our longs are AI retail favourites (chance 10%); "
        f"an AI sell-off of $50B would cost us {f['fv3_sell_50']}.",
        "Next: run the final book at any size from $10M to $5B; watch the AI crowd as a risk "
        "factor; "
        "use the stock-picking model only on cheap-to-trade stocks.",
    ]
    subtitle = (
        "January 2021 to August 2026. Total returns after all modelled costs, position P&L "
        "before costs"
    )
    page = D.slide(
        deck, "What Worked, What Did Not, and What Comes Next", subtitle, section("Discussion")
    )
    items = [tuple(line.split(": ", 1)) if ": " in line else ("", line) for line in lines]
    items = [(lead + "." if lead else "", rest[:1].upper() + rest[1:]) for lead, rest in items]
    D.text(page, items, D.LEFT, Inches(1.65), Inches(5.75), Inches(5.3), 10.5, gap=5)
    D.chart(page, "pnl_contributors", Inches(6.55), Inches(1.7), Inches(6.18), Inches(5.2))
    finish(page)


# Appendix: each fix's own pre-2021 test, then the edge tests.
def evidence(deck, f: dict, m: dict) -> None:
    verdicts = load(V2 / "evidence.json").get("verdicts", {})
    rows = ["| Fix | Its own pre-2021 test |", "|---|---|"]
    rows += [
        f"| {k} | {'pass' if v.get('passed') else 'fail'} |"
        for k, v in verdicts.items()
        if k != "v2"
    ]
    subtitle = (
        "Fixes judged on 2016-2020 only. Edge tests on January 2021 to August 2026, after costs"
    )
    heading = "Every Fix on Its Own Test, and Is the Edge Real?"
    page = D.slide(deck, heading, subtitle, section("Appendix", True))
    D.table(page, "\n".join(rows), Inches(0.6), Inches(1.6), Inches(4.5))
    edge = [
        "| Test | Result |",
        "|---|---|",
        f"| Deflated Sharpe after all {f['trials']} trials, first and final book | "
        f"{f['dsr_v1']} and {f['dsr_fv3']} |",
        f"| First book's alpha t on FF5 and momentum, Newey-West 3 to 11 lags | {f['nw3']} to "
        f"{f['nw11']} |",
        f"| First book's IC needed to beat cash plus 4%, and realized | {f['ic_star']} and "
        f"{f['ic_real']} |",
        f"| Our spreads against TAQ, large and micro caps | {f['taq_large']}x and "
        f"{f['taq_micro']}x |",
    ]
    final = load(V2 / "evidence.json").get("v3", {}).get("final", {}).get("name")
    replay = load(V2 / "nautilus" / "check.json").get(f"{final}@1e7")
    if replay:
        edge.append(
            "| NautilusTrader replays the final book in whole shares, 2021 to 2025, gap a month "
            f"and in total | {replay['mean_abs_gap']:.4%} and {replay['cumulative_gap']:+.2%}, "
            "limit 0.10% |"
        )
    D.table(page, "\n".join(edge), Inches(5.6), Inches(1.6), Inches(7.13))
    finish(page)


# Deck 1's prompts, then the finals agents' own, as the brief asks.
def prompts(deck, f: dict, m: dict) -> None:
    from fiam.agents import crowd_lab, rules_reader

    subtitle = (
        "The system prompt each local model ran with, punctuation lightly edited. Exact text in "
        "prompts/agents/"
    )
    page = D.slide(deck, "The Agents' Prompts", subtitle, "Appendix")
    left, right = D.prompts()
    D.text(page, left, D.LEFT, Inches(1.6), D.HALF, Inches(5.4), 8.5, gap=3.5)
    D.text(page, right, D.RIGHT, Inches(1.6), D.HALF, Inches(5.4), 8.5, gap=3.5)
    finish(page)
    retail = crowd_lab.PROMPT.read_text(encoding="utf-8").split("\n## ")
    copycat = crowd_lab.COPYCAT.read_text(encoding="utf-8").strip()
    subtitle = (
        "{persona}, {picks} and {strategy} are filled in per agent. Exact text in prompts/agents/ "
        "and fiam/agents/rules_reader.py"
    )
    page = D.slide(deck, "The AI Crowd Lab's and the Rules Reader's Prompts", subtitle, "Appendix")
    left = [("AI retail investor, Gemma 4 26B-A4B.", "")]
    left += [line for line in retail[0].split("\n") if line.strip()]
    left += [(s.split("\n", 1)[0] + ".", s.split("\n", 1)[1].strip()) for s in retail[1:]]
    right = [("AI copycat fund, Gemma 4 26B-A4B.", "")]
    right += [line for line in copycat.split("\n") if line.strip()]
    right += [("Rules reader, Gemma 4 26B-A4B.", rules_reader.SYSTEM)]
    D.text(page, left, D.LEFT, Inches(1.6), D.HALF, Inches(5.4), 9, gap=3.5)
    D.text(page, right, D.RIGHT, Inches(1.6), D.HALF, Inches(5.4), 9, gap=3.5)
    finish(page)


# The critic's tools as the model got them, the data and how to run it.
def tools(deck, f: dict, m: dict) -> None:
    with submitted():
        listed = R.load("exhibits/agents/tools.json")
    subtitle = (
        "The devil's advocate's tools as the model received them, and what the portfolio was "
        "built from"
    )
    page = D.slide(deck, "Tools, Data and How to Run It", subtitle, "Appendix")
    rows = ["| Tool | What It Returns | Arguments |", "|---|---|---|"]
    D.VERBATIM.update(t["name"] for t in listed.get("critic", []))
    for item in listed.get("critic", []):
        args = item["parameters"].get("properties", {})
        D.VERBATIM.update(args)
        shown = ", ".join(f"{k} ({v.get('type', '')})" for k, v in args.items())
        rows.append(f"| {item['name']} | {item['description']} | {shown or 'none'} |")
    below = D.table(page, "\n".join(rows), D.LEFT, Inches(1.65), D.WIDTH, 10, 0.28)
    text = [
        (
            "Data.",
            "The organizer's stock panel and 8-K filings, plus WRDS (CRSP delistings, daily "
            "prices and quotes, Compustat report dates and short interest, I/B/E/S), FRED's "
            "T-bill and S&P 500, Kenneth French's factors and SEC EDGAR names, joined on PERMNO "
            "and month.",
        ),
        (
            "Run it.",
            "pip install uv, then uv run MAIN.py. On its own it checks our files against "
            "the brief. With the organizer's three data files it rebuilds the portfolio. Every AI "
            "answer is cached with its prompt and model hash, so the agents replay without a GPU.",
        ),
    ]
    D.text(page, text, D.LEFT, below + Inches(0.3), D.WIDTH, Inches(2.4), 13, gap=10)
    finish(page)


# The live talk, which ends on its closing slide, and the printed deck.
# fmt: off
LIVE = [title, hook, ideas, journey, ml, desk, results, growth, neutral, edge, monkeys, ladder,
        capacity, final, risk_slide, exposures_slide, losses, turn, meet, swarm, film, copycats,
        trades, receipt, today, launch, closing]
PAPER = [summary, strategy, method, committee, discussion, final, edge_page, losses, risk_page,
         crowd_page, receipt, evidence, prompts, tools]
# fmt: on


# Builds a deck from its slides, swaps in our footer, returns the count.
def build(parts: list, target: Path, f: dict, m: dict) -> int:
    deck = Presentation()
    deck.slide_width, deck.slide_height = D.WIDE, D.TALL
    for part in parts:
        part(deck, f, m)
    for page in deck.slides:
        for item in page.shapes:
            runs = (
                [r for p in item.text_frame.paragraphs for r in p.runs]
                if item.has_text_frame
                else []
            )
            for run in runs:
                if run.text == FOOTERS[0]:
                    run.text = FOOTERS[1]
    target.parent.mkdir(exist_ok=True)
    deck.save(target)
    return len(deck.slides)


# The PDF by LibreOffice, else PowerPoint on Windows, else None.
def pdf(path: Path) -> Path | None:
    found = [
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    ]
    office = shutil.which("soffice") or next((str(p) for p in found if p.exists()), None)
    target = path.with_suffix(".pdf")
    if office:
        command = [
            office,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(path.parent),
            str(path),
        ]
    elif sys.platform != "win32":
        return None
    else:
        script = (
            f"$p = New-Object -ComObject PowerPoint.Application; $d = $p.Presentations.Open("
            f"'{path}', $true, $false, $false); $d.SaveAs('{target}', 32); $d.Close(); $p.Quit()"
        )
        command = ["powershell", "-NoProfile", "-Command", script]
    try:
        subprocess.run(command, check=True)
    except OSError, subprocess.CalledProcessError:
        return None
    return target if target.exists() else None


# Builds the live deck and the printed one, and the PDF with --pdf.
def main(args: list[str]) -> None:
    global MODE
    if "finals" not in paths.OUT_ROOT.as_posix():
        raise SystemExit("The finals deck reads final v3's pack:  uv run run.py --finals-pack")
    D.FIGURES = ROOT / paths.SUBMISSION / "figures"
    f, m = qa.facts(), R.load("metrics.json")
    MODE = "live"
    print(f"wrote {OUT}: {build(LIVE, OUT, f, m)} slides")
    MODE = "print"
    pages = build(PAPER, PRINTED, f, m)
    print(f"wrote {PRINTED}: {pages} pages, 8 for the brief and {pages - 8} in the appendix")
    if pages > 18:
        print("the brief allows an appendix of 10 pages at most")
    if "--pdf" in args:
        final, made = OUT.with_suffix(".pdf"), pdf(PRINTED)
        if made:
            made.replace(final)
            print(f"wrote {final}")
        else:
            print("PDF skipped: install LibreOffice (or PowerPoint on Windows) for it.")


if __name__ == "__main__":
    main(sys.argv[1:])
