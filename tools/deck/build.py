# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "python-pptx==1.0.2",
#     "pillow==12.3.0",
#     "lxml==6.1.3",
#     "imageio-ffmpeg==0.7.0",
# ]
# [tool.uv]
# exclude-newer = "2026-10-30T00:00:00Z"
# ///
# Rebuilds finals/QUARCC_Purple_Finals.pptx and .pdf from deck.html.
import itertools
import json
import shutil
import subprocess
import time
import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import imageio_ffmpeg
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BUILD = HERE / "build"
OUT = ROOT / "finals" / "QUARCC_Purple_Finals"
# One set date inside the files, so a rebuild gives the same bytes.
STAMP = datetime(2026, 10, 29, tzinfo=UTC)
# A 1920 pixel page on the 13.333 inch slide.
PX = 6350
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
CLIP_ALT = "30 AI agents gather into one shared bet"


# Pixels on the 1920 wide page as EMU on the slide.
def emu(v: float) -> Emu:
    return Emu(round(v * PX))


# One entrance: made visible at once, then a left wipe or a fade.
def entrance(kind: str, spid: int, dur: int, ids) -> str:
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    flt = "fade" if kind == "fade" else "wipe(left)"
    return (
        f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst>'
        f'<p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}<p:attrNameLst>'
        "<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
        '<p:to><p:strVal val="visible"/></p:to></p:set>'
        f'<p:animEffect transition="in" filter="{flt}"><p:cBhvr>'
        f'<p:cTn id="{next(ids)}" dur="{dur}"/>{tgt}</p:cBhvr></p:animEffect>'
    )


# The preset PowerPoint shows for each kind: wipe from left or fade.
def preset(kind: str) -> str:
    if kind == "wipe":
        return 'presetID="22" presetClass="entr" presetSubtype="8"'
    return 'presetID="10" presetClass="entr" presetSubtype="0"'


# The slide timeline, every effect timed from the slide's start.
def timeline(steps: list, media: list) -> str:
    ids = itertools.count(4)
    pars = []
    for spid, kind, start, dur in steps:
        outer, inner = next(ids), next(ids)
        body = entrance(kind, spid, dur, ids)
        pars.append(
            f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="{start}"/>'
            f'</p:stCondLst><p:childTnLst><p:par><p:cTn id="{inner}" {preset(kind)} '
            'fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f"<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
        )
    # Each clip plays from the start as the slide opens, muted.
    for spid in media:
        outer, inner, call = next(ids), next(ids), next(ids)
        pars.append(
            f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst><p:childTnLst><p:par><p:cTn id="{inner}" presetID="1" '
            'presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" '
            f'cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{call}" dur="indefinite" fill="hold"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst>'
            "</p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
        )
    nodes = ""
    for spid in media:
        nodes += (
            f'<p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="{next(ids)}" fill="hold" '
            'display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst>'
            '<p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
            f'</p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode>'
            "</p:video>"
        )
    cond = '<p:cond evt="{}" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
    return (
        f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" '
        'restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
        '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" '
        'fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0">'
        f'<p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst>'
        f"</p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst>{cond.format('onPrev')}"
        f"</p:prevCondLst><p:nextCondLst>{cond.format('onNext')}</p:nextCondLst></p:seq>"
        f"{nodes}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
    )


# The timeline placed after the colour map, as the schema orders it.
def finish(slide, steps: list, media: list) -> None:
    sld = slide._element
    for tag in ("p:transition", "p:timing"):
        for old in sld.findall(qn(tag)):
            sld.remove(old)
    at = list(sld).index(sld.find(qn("p:clrMapOvr"))) + 1
    sld.insert(at, etree.fromstring(timeline(steps, media)))


# A full slide frame as a high-quality jpeg, much smaller to send.
def jpg(png: Path) -> str:
    out = png.with_suffix(".jpg")
    Image.open(png).convert("RGB").save(out, "JPEG", quality=92, optimize=True)
    return str(out)


# The crowd clip at 30 frames a second, its last frame the poster.
def encode_clip() -> None:
    frames = sorted((BUILD / "frames").glob("f_*.png"))
    subprocess.run(
        [
            imageio_ffmpeg.get_ffmpeg_exe(),
            "-y",
            "-loglevel",
            "error",
            "-framerate",
            "30",
            "-i",
            str(BUILD / "frames" / "f_%04d.png"),
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            "-vf",
            "scale=trunc(iw/2)*2:trunc(ih/2)*2",
            "-crf",
            "18",
            "-movflags",
            "+faststart",
            "-map_metadata",
            "-1",
            "-fflags",
            "+bitexact",
            "-flags:v",
            "+bitexact",
            str(BUILD / "crowd.mp4"),
        ],
        check=True,
    )
    shutil.copy(frames[-1], BUILD / "crowd_poster.png")


# One slide: the base frame, its moving parts and the clip on top.
def add_slide(prs, m: dict, size: tuple) -> int:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if not m.get("base"):
        s.shapes.add_picture(jpg(BUILD / m["final"]), 0, 0, *size)
        return 0
    s.shapes.add_picture(jpg(BUILD / m["base"]), 0, 0, *size)
    full = Image.open(BUILD / m["final"])
    steps, media = [], []
    for i, r in enumerate(m["regions"]):
        box = (r["x"] * 2, r["y"] * 2, (r["x"] + r["w"]) * 2, (r["y"] + r["h"]) * 2)
        part = BUILD / f"ov_{m['k'] + 1:02d}_{i}.png"
        full.crop(tuple(round(v) for v in box)).save(part)
        pic = s.shapes.add_picture(str(part), emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
        steps.append((pic.shape_id, r.get("kind", "wipe"), 300 + i * 250, 1100))
    if m.get("clip"):
        c = m["clip"]
        mv = s.shapes.add_movie(
            str(BUILD / "crowd.mp4"),
            emu(c["x"]),
            emu(c["y"]),
            emu(c["w"]),
            emu(c["h"]),
            poster_frame_image=str(BUILD / "crowd_poster.png"),
            mime_type="video/mp4",
        )
        mv._element.nvPicPr.cNvPr.set("descr", CLIP_ALT)
        media.append(mv.shape_id)
    finish(s, steps, media)
    return len(steps) + len(media)


# The pptx with set dates inside it and on every zip entry.
def save_pptx(prs, path: Path) -> None:
    core = prs.core_properties
    core.title = "QUARCC Purple Finals"
    core.author = core.last_modified_by = "QUARCC Purple"
    core.created = core.modified = STAMP.replace(tzinfo=None)
    core.revision = 1
    raw = BytesIO()
    prs.save(raw)
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(path, "w") as out:
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=STAMP.timetuple()[:6])
            entry.compress_type, entry.create_system = zipfile.ZIP_DEFLATED, 3
            out.writestr(entry, src.read(info.filename))


# Frames from Node, the clip, then the pptx and the PDF.
def main() -> None:
    subprocess.run(["node", str(HERE / "frames.js")], check=True)
    encode_clip()
    meta = json.loads((BUILD / "meta.json").read_text(encoding="utf-8"))["meta"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    size = (prs.slide_width, prs.slide_height)
    effects = sum(add_slide(prs, m, size) for m in meta)
    OUT.parent.mkdir(exist_ok=True)
    save_pptx(prs, OUT.with_suffix(".pptx"))
    pages = [Image.open(BUILD / m["final"]).convert("RGB") for m in meta]
    stamp = time.gmtime(STAMP.timestamp())
    pages[0].save(
        OUT.with_suffix(".pdf"),
        save_all=True,
        append_images=pages[1:],
        resolution=288,
        title="QUARCC Purple Finals",
        creationDate=stamp,
        modDate=stamp,
    )
    print(f"{len(prs.slides)} slides, {effects} animations, written to {OUT.parent}")


if __name__ == "__main__":
    main()
